#!/usr/bin/env python3
"""Export the minimal far-view scale cards from the reviewed workbook."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

import openpyxl


def sheet_records(workbook, name: str) -> list[dict]:
    sheet = workbook[name]
    headers = list(next(sheet.iter_rows(min_row=3, max_row=3, values_only=True)))
    return [
        dict(zip(headers, row))
        for row in sheet.iter_rows(min_row=4, values_only=True)
        if row[0]
    ]


def number(value) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def build_cards(workbook) -> list[dict]:
    cards = []

    for row in sheet_records(workbook, "坊功能表"):
        cards.append(
            {
                "id": row["fang_key*"],
                "type": "Fang",
                "name": row["name*"],
                "scale": (
                    f'{number(row["east_west_width_m"])} × '
                    f'{number(row["north_south_depth_m"])} 米'
                ),
                "symbol": "坊",
            }
        )

    for row in sheet_records(workbook, "道路分级表"):
        normalized = row.get("normalized_width_m")
        original = row.get("original_width_value")
        unit = row.get("original_width_unit")
        if original is not None and unit:
            scale = f"宽 {number(original)} {unit}"
        elif normalized is not None:
            scale = f"宽 {number(normalized)} 米"
        else:
            scale = "尺度暂无"
        cards.append(
            {
                "id": row["road_key*"],
                "type": "Road",
                "name": row["name*"],
                "scale": scale,
                "symbol": "街",
            }
        )

    return cards


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    workbook = openpyxl.load_workbook(args.workbook, read_only=True, data_only=True)
    cards = build_cards(workbook)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    payload = json.dumps(cards, ensure_ascii=False, indent=2)
    args.output.write_text(f"window.CHANGAN_SCALE_CARDS = {payload};\n", encoding="utf-8")
    print(f"Exported {len(cards)} cards to {args.output}")


if __name__ == "__main__":
    main()
